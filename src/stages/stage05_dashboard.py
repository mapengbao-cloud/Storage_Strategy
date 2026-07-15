"""Stage 05: Dashboard / statistics table update (收益统计表).

Reads review files from earlier stages and writes computed summary values
into the master statistics table 山东夏津储能收益统计表.xlsx.

Data source rule (critical):
- 日前列 (B-R) → compute_revenue_from_96point() from 日前 review file
- 实时列 (S-AI) → compute_revenue_from_96point() from 实时 review file
- 日结算列 (AJ-BA) → compute_revenue_from_settlement() from 日结算 review file

Usage (from pipeline):
    from src.stages.stage05_dashboard import run
    run(['0522', '0523', '0524'])

Usage (standalone):
    python -m src.stages.stage05_dashboard --day-ahead 0522-0531
"""

import re
import sys
import shutil
import argparse
import openpyxl
from pathlib import Path
from collections import defaultdict
from datetime import datetime, date

from src.data.readers import read_trading_result_xls, read_settlement_base_values
from src.data.models import DailyRevenue
from src.business.revenue import (
    RevenueInput,
    compute_revenue_from_96point,
    compute_revenue_from_settlement,
)
from src.business.capacity import compute_J_val_from_workbook
from src.utils.numerics import safe_float, round_value, PCT_COLUMNS
from src.utils.date_utils import mmdd_to_date, date_to_iso
from src.config import get_month_column


# ── Column mappings ──────────────────────────────────────────────

def _linear_mapping(start_col: int, count: int = 17) -> dict:
    """Map source cols 1-17 to target cols starting at start_col."""
    return {i + 1: start_col + i for i in range(count)}

DAY_AHEAD_MAP = _linear_mapping(2)    # B-R (col 2-18)
REAL_TIME_MAP = _linear_mapping(19)   # S-AI (col 19-35)

# 日结算: A-O → AJ-AX (36-50), P → AZ (52), Q → BA (53), AY(51)=O6
SETTLEMENT_MAP = {i: 36 + (i - 1) for i in range(1, 16)}  # A-O
SETTLEMENT_MAP[16] = 52  # P → AZ
SETTLEMENT_MAP[17] = 53  # Q → BA
SETTLEMENT_AY_COL = 51   # O6 column


# ── Target row lookup ────────────────────────────────────────────

def find_date_row(ws, month: int, day: int) -> int | None:
    """Find the row in the target sheet matching month/day.

    Handles both datetime objects and Excel serial numbers in column A.
    """
    excel_epoch = datetime(1899, 12, 30)
    for year in [2026, 2025]:
        try:
            target_serial = (datetime(year, month, day) - excel_epoch).days
        except ValueError:
            continue
        for row in range(4, ws.max_row + 1):
            cell_val = ws.cell(row=row, column=1).value
            if isinstance(cell_val, datetime):
                if (cell_val.year == year
                        and cell_val.month == month
                        and cell_val.day == day):
                    return row
            elif isinstance(cell_val, (int, float)) and cell_val > 40000:
                if int(cell_val) == target_serial:
                    return row
    return None


# ── Source file discovery ────────────────────────────────────────

def _parse_date_from_name(filename: str) -> tuple[int, int] | None:
    """Extract (month, day) from MMDD-prefixed filename."""
    m = re.match(r'^(\d{2})(\d{2})', filename)
    return (int(m.group(1)), int(m.group(2))) if m else None


def discover_source_files(source_dir: str) -> dict:
    """Scan directory for review files, group by date.

    Returns:
        {(month, day): {'day_ahead': str, 'real_time': str, 'settlement': str}}
    """
    source_files = defaultdict(dict)
    src_path = Path(source_dir)

    for fpath in src_path.iterdir():
        fname = fpath.name
        if not fname.endswith('.xlsx') or fname.startswith('~$'):
            continue
        if '统计表' in fname or 'template' in fname.lower():
            continue

        date_info = _parse_date_from_name(fname)
        if date_info is None:
            continue

        md = date_info
        if '日前' in fname and '机组组合' in fname:
            source_files[md]['day_ahead'] = str(fpath)
        elif '实时' in fname and '机组组合' in fname:
            source_files[md]['real_time'] = str(fpath)
        elif '日结算' in fname and '收益复盘' in fname:
            source_files[md]['settlement'] = str(fpath)

    return dict(source_files)


# ── Main run function ────────────────────────────────────────────

def run(
    date_list: list[str],
    source_dir: str = "output",
    output_dir: str = "output/reports",
    target_file: str = "山东夏津储能收益统计表.xlsx",
    template_path: str | None = None,
) -> Path:
    """Update the master statistics table for given dates.

    Args:
        date_list: List of MMDD strings (e.g. ['0522', '0523']).
        source_dir: Directory containing stage 02/03/04 output files.
        output_dir: Directory for the statistics table.
        target_file: Name of the statistics table file.
        template_path: Optional path to statistics template.
            Defaults to assets/templates/山东夏津储能收益统计表.xlsx.

    Returns:
        Path to the saved statistics table.
    """
    # Resolve template
    if template_path:
        tpl = Path(template_path)
    else:
        tpl = Path("assets/templates") / target_file
    if not tpl.exists():
        raise FileNotFoundError(f"Template not found: {tpl}")

    out_dir = Path(output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / target_file

    # Use previous output as base (incremental), fall back to template
    base = out_path if out_path.exists() else tpl
    print(f"Base file: {base}")

    try:
        shutil.copy2(base, out_path)
    except PermissionError:
        ts = datetime.now().strftime('%Y%m%d_%H%M%S')
        stem, ext = target_file.rsplit('.', 1)
        out_path = out_dir / f"{stem}_{ts}.{ext}"
        shutil.copy2(base, out_path)

    # Discover source files
    source_files = discover_source_files(source_dir)

    # Determine month for capacity coefficient column
    # Use the first date in the list
    first_date = date_list[0]
    month, _ = int(first_date[:2]), int(first_date[2:])

    # Open target workbook
    wb = openpyxl.load_workbook(out_path)
    ws = wb['润津']

    for date_str in date_list:
        month_num, day_num = int(date_str[:2]), int(date_str[2:])
        target_row = find_date_row(ws, month_num, day_num)
        if target_row is None:
            print(f"  WARNING: {date_str} not found in target, skipping")
            continue

        files = source_files.get((month_num, day_num), {})
        if not files:
            print(f"  WARNING: No source files for {date_str}, skipping")
            continue

        print(f"Processing {month_num:02d}/{day_num:02d} → row {target_row}")

        # ── 日前 ──
        if 'day_ahead' in files:
            print(f"  - 日前: {Path(files['day_ahead']).name}")
            _write_day_ahead(ws, target_row, files['day_ahead'])

        # ── 实时 ──
        if 'real_time' in files:
            print(f"  - 实时: {Path(files['real_time']).name}")
            _write_real_time(ws, target_row, files['real_time'])

        # ── 日结算 ──
        if 'settlement' in files:
            print(f"  - 日结算: {Path(files['settlement']).name}")
            _write_settlement(ws, target_row, files['settlement'])

    # Copy number_format from previous row for all written rows
    _copy_formats_from_prev_row(ws, month_num, day_num, date_list)

    save_and_close(wb, out_path)
    print(f"\nSaved: {out_path}")
    return out_path


def save_and_close(wb: openpyxl.Workbook, path: Path) -> None:
    """Save workbook and close."""
    wb.save(str(path))
    wb.close()


# ── Section writers ──────────────────────────────────────────────

def _write_day_ahead(ws, target_row: int, src_path: str):
    """Write 日前 section (B-R, col 2-18) from day-ahead review file."""
    revenue = _compute_from_trading_review(src_path)
    _write_values(ws, target_row, revenue.to_list(), DAY_AHEAD_MAP)


def _write_real_time(ws, target_row: int, src_path: str):
    """Write 实时 section (S-AI, col 19-35) from real-time review file."""
    revenue = _compute_from_trading_review(src_path)
    _write_values(ws, target_row, revenue.to_list(), REAL_TIME_MAP)


def _write_settlement(ws, target_row: int, src_path: str):
    """Write 日结算 section (AJ-BA) from settlement review file."""
    revenue = _compute_from_settlement_file(src_path)
    values = revenue.to_list()
    _write_values(ws, target_row, values, SETTLEMENT_MAP)

    # O6 → AY column (51)
    cell = ws.cell(row=target_row, column=SETTLEMENT_AY_COL)
    cell.value = round_value(revenue.O6)


def _write_values(ws, target_row: int, values: list[float],
                  mapping: dict[int, int]):
    """Write source values to target columns using mapping."""
    for src_col, val in enumerate(values, 1):
        target_col = mapping.get(src_col)
        if target_col is None:
            continue
        cell = ws.cell(row=target_row, column=target_col)
        cell.value = round_value(val, target_col)


def _copy_formats_from_prev_row(
    ws, month_num: int, day_num: int, date_list: list[str]
):
    """Copy number_format from previous row for newly written rows."""
    for date_str in date_list:
        m, d = int(date_str[:2]), int(date_str[2:])
        target_row = find_date_row(ws, m, d)
        if target_row is None or target_row < 5:
            continue
        prev_row = target_row - 1
        # Copy formats from cols 2-53 (B-BA)
        for col in range(2, 54):
            prev_cell = ws.cell(row=prev_row, column=col)
            curr_cell = ws.cell(row=target_row, column=col)
            if prev_cell.number_format and prev_cell.number_format != 'General':
                curr_cell.number_format = prev_cell.number_format


# ── Revenue computation ──────────────────────────────────────────

def _compute_from_trading_review(src_path: str) -> DailyRevenue:
    """Compute DailyRevenue from a 日前/实时 review file.

    Reads 报价及预中标 raw data, 充放测算 parameters, and 容量分摊系数.
    """
    cd_data = read_trading_result_xls(Path(src_path))

    wb = openpyxl.load_workbook(src_path, data_only=False)
    sn0 = wb.sheetnames[0]  # 充放测算
    sn_cap = next(s for s in wb.sheetnames if '容量分摊' in s)
    ws0 = wb[sn0]
    ws_cap = wb[sn_cap]

    # Extract date from filename to determine month
    fname = Path(src_path).name
    m = re.match(r'^(\d{2})(\d{2})', fname)
    month = int(m.group(1)) if m else 5

    I8 = safe_float(ws0.cell(row=8, column=9).value)
    I9 = safe_float(ws0.cell(row=9, column=9).value)
    I12 = safe_float(ws0.cell(row=12, column=9).value)
    I13 = safe_float(ws0.cell(row=13, column=9).value)
    I14 = safe_float(ws0.cell(row=14, column=9).value)

    # Compute J_val from capacity allocation sheet
    price_ws = _find_price_sheet(wb)
    J_val = compute_J_val_from_workbook(ws_cap, price_ws, month)

    wb.close()
    return compute_revenue_from_96point(cd_data, I8, I9, I12, I13, I14, J_val)


def _compute_from_settlement_file(src_path: str) -> DailyRevenue:
    """Compute DailyRevenue from a 日结算 review file.

    Reads settlement base values from 充电/放电日清算费用 sheets
    and parameters from 充放测算.
    """
    wb = openpyxl.load_workbook(src_path, data_only=True)
    ws_cf = wb['充放测算']
    ws_charge = wb['充电日清算费用']
    ws_discharge = wb['放电日清算费用']

    # Base values
    A = safe_float(ws_charge.cell(row=29, column=29).value)      # AC29
    B = -safe_float(ws_charge.cell(row=29, column=28).value)     # -AB29
    C = -safe_float(ws_charge.cell(row=29, column=30).value)     # -AD29
    L = safe_float(ws_discharge.cell(row=101, column=16).value)  # P101
    M_val = safe_float(ws_discharge.cell(row=101, column=41).value)  # AO101
    N_val = safe_float(ws_discharge.cell(row=101, column=17).value)  # Q101

    # Parameters
    I8 = safe_float(ws_cf.cell(row=8, column=9).value)
    I9 = safe_float(ws_cf.cell(row=9, column=9).value)
    I12 = safe_float(ws_cf.cell(row=12, column=9).value)
    I13 = safe_float(ws_cf.cell(row=13, column=9).value)
    I14 = safe_float(ws_cf.cell(row=14, column=9).value)
    J_val = safe_float(ws_cf.cell(row=4, column=10).value)

    wb.close()
    return compute_revenue_from_settlement(
        A, B, C, L, M_val, N_val, I8, I9, I12, I13, I14, J_val
    )


def _find_price_sheet(wb) -> any:
    """Find the 报价及预中标 sheet in a workbook."""
    for sn in wb.sheetnames:
        if '报价' in sn and '预中标' in sn:
            return wb[sn]
    raise ValueError("No 报价及预中标 sheet found")


# ── CLI ──────────────────────────────────────────────────────────

def _parse_range(arg: str) -> list[str]:
    """Parse 'MMDD-MMDD' or 'MMDD' into list of MMDD strings."""
    from src.utils.date_utils import expand_mmdd_range
    dates = expand_mmdd_range(arg)
    return [d.strftime("%m%d") for d in dates]


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description='Update energy storage revenue statistics table.')
    parser.add_argument('--day-ahead', type=str,
                        help='Date range for 日前 data (MMDD-MMDD)')
    parser.add_argument('--real-time', type=str,
                        help='Date range for 实时 data (MMDD-MMDD)')
    parser.add_argument('--settlement', type=str,
                        help='Date range for 日结算 data (MMDD-MMDD)')
    parser.add_argument('--source-dir', type=str, default='output',
                        help='Directory with stage 02/03/04 output')
    args = parser.parse_args()

    # Collect all dates
    all_dates = set()
    for r in [args.day_ahead, args.real_time, args.settlement]:
        if r:
            all_dates.update(_parse_range(r))

    if not all_dates:
        print("Please specify at least one date range")
        sys.exit(1)

    date_list = sorted(all_dates)
    print(f"Processing {len(date_list)} dates: {date_list[0]}...{date_list[-1]}")
    run(date_list, source_dir=args.source_dir)