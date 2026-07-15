"""Stage 04: Daily settlement review (日结算收益复盘).

Consolidates charge/discharge settlement statements with real-time trading
parameters into a daily review workbook.

Requires three inputs per date:
1. Charge settlement statement (~11KB)
2. Discharge settlement statement (~21KB)
3. Real-time trading review (from Stage 03)

Usage (from pipeline):
    from src.stages.stage04_settlement import run
    run('0605')

Template: assets/templates/日结算收益复盘.xlsx
Output:  output/YYYY-MM-DD/MMDD-日结算收益复盘.xlsx
"""

import sys
from pathlib import Path
from typing import Optional

import openpyxl

from src.data.writers import (
    copy_template,
    write_settlement_review,
    write_review_parameters,
    save_and_close,
)
from src.business.capacity import compute_J_val_from_workbook
from src.utils.date_utils import mmdd_to_date, date_to_iso
from src.config import get_month_column


def run(
    date_mmdd: str,
    charge_stmt_path: Optional[str] = None,
    discharge_stmt_path: Optional[str] = None,
    rt_review_path: Optional[str] = None,
    data_root: str = "data",
    output_root: str = "output",
    template_root: str = "assets/templates",
) -> Path:
    """Generate daily settlement review for one date.

    Args:
        date_mmdd: MMDD date string (e.g. '0605').
        charge_stmt_path: Optional explicit path to charge settlement.
        discharge_stmt_path: Optional explicit path to discharge settlement.
        rt_review_path: Optional explicit path to RT review file.
        data_root: Root directory for source data.
        output_root: Root directory for output files.
        template_root: Root directory for templates.

    Returns:
        Path to the generated output file.
    """
    d = mmdd_to_date(date_mmdd)
    date_iso = date_to_iso(d)
    month = date_iso.month if hasattr(date_iso, 'month') else d.month
    out_name = f"{date_mmdd}-日结算收益复盘.xlsx"

    data_dir = Path(data_root) / "raw" / date_iso

    # Resolve source files
    def _resolve(path, pattern):
        if path:
            return Path(path)
        candidates = list(data_dir.glob(pattern))
        if not candidates:
            raise FileNotFoundError(
                f"No file matching '{pattern}' for {date_mmdd} in {data_dir}"
            )
        return candidates[0]

    charge_path = _resolve(charge_stmt_path, "*结算单-充电*")
    discharge_path = _resolve(discharge_stmt_path, "*结算单-放电*")
    rt_path = _resolve(rt_review_path, "*实时机组组合收益复盘*")

    # Open template
    tpl = Path(template_root) / "日结算收益复盘.xlsx"
    if not tpl.exists():
        raise FileNotFoundError(f"Template not found: {tpl}")

    out_dir = Path(output_root) / date_iso
    out_path = out_dir / out_name
    wb_out = copy_template(tpl, out_path)

    # Load source workbooks
    wb_charge = openpyxl.load_workbook(charge_path, data_only=True)
    wb_discharge = openpyxl.load_workbook(discharge_path, data_only=True)
    wb_rt = openpyxl.load_workbook(rt_path, data_only=True)

    # Copy settlement data
    write_settlement_review(wb_out, wb_charge, wb_discharge)

    # Compute J4 (capacity allocation coefficient) from RT review
    rt_cap_ws = wb_rt["容量分摊系数"]
    rt_price_ws = wb_rt["报价及预中标"]
    j4_val = compute_J_val_from_workbook(rt_cap_ws, rt_price_ws, month)

    # Write parameters to 充放测算
    rt_cf = wb_rt["充放测算"]
    out_cf = wb_out["充放测算"]
    write_review_parameters(out_cf, rt_cf, j4_value=j4_val)

    # Close source workbooks
    for wb in [wb_charge, wb_discharge, wb_rt]:
        wb.close()

    save_and_close(wb_out, out_path)
    print(f"Saved: {out_path}")
    return out_path


# ── CLI entry point ──────────────────────────────────────────────

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    run(sys.argv[1])