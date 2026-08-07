"""Generate 实时/日前出清收益测算 per 收益测算工作流程.md.

Step 1: Query 天机 MySQL directly for 96-point data
Step 2: Copy 按月 template → output
Step 3: Write 报价及预中标 (H=timestamp, J/K=price, N=power)
Step 4: COM refresh → openpyxl read Row 4 → SQLite ingest
"""
import json, os, shutil, sqlite3, sys

import openpyxl
import pymysql
import pythoncom
import win32com.client

PROJECT = r'E:\DataWork\Storage_Strategy'
TEMPLATES = os.path.join(PROJECT, 'assets', 'templates', '收益测算')
DB_PATH = os.path.join(PROJECT, 'data', 'cache', 'local.db')

RT_OUT = os.path.join(PROJECT, 'output', '实时出清收益测算')
DA_OUT = os.path.join(PROJECT, 'output', '日前出清收益测算')
os.makedirs(RT_OUT, exist_ok=True)
os.makedirs(DA_OUT, exist_ok=True)

TIMES = [f'{h:02d}:{m:02d}' for h in range(24) for m in (0, 15, 30, 45)]


def _get_template(month_int: int, mode: str) -> str:
    """Select the correct monthly template."""
    if mode == 'realtime':
        name = f'实时出清收益测算-{month_int}月.xlsx'
    else:
        name = f'日前出清收益测算-{month_int}月.xlsx'
    path = os.path.join(TEMPLATES, name)
    if not os.path.exists(path):
        raise FileNotFoundError(f'Template not found: {path}')
    return path


def get_conn():
    return pymysql.connect(
        host='rm-2zej7q7186wi4eds5no.mysql.rds.aliyuncs.com', port=3306,
        user='pengyiqiang', password='pengyiqiang123', database='tianrun_new',
        charset='utf8mb4', connect_timeout=10, read_timeout=30)


def com_refresh(path):
    pythoncom.CoInitialize()
    try:
        excel = win32com.client.Dispatch('Excel.Application')
        excel.Visible = False
        excel.DisplayAlerts = False
        wb = excel.Workbooks.Open(os.path.abspath(path))
        wb.RefreshAll()
        excel.CalculateUntilAsyncQueriesDone()
        wb.Save()
        wb.Close()
        excel.Quit()
    except Exception:
        pass


def safe_float(v):
    if v is None:
        return None
    try:
        return float(v)
    except (ValueError, TypeError):
        return None


def _is_merged(cell):
    return type(cell).__name__ == 'MergedCell'


def generate_one(mmdd, mode='realtime'):
    """Generate one day's clearing revenue Excel.

    mode: 'realtime' or 'dayahead'
    """
    date_iso = f'2026-{mmdd[:2]}-{mmdd[2:]}'

    if mode == 'realtime':
        template = _get_template(int(mmdd[:2]), 'realtime')
        out_dir = RT_OUT
        out_name = f'{mmdd}-实时出清收益测算.xlsx'
        sql = f"""SELECT date, time_point, power, price
FROM shandong_px_realtime_clearing_results_query
WHERE date = '{date_iso}' AND member_id = 'b9e64e64a713458eba94c9af05c0a757'
ORDER BY time_point"""
        table_name = '实时出清收益测算'
        raw_table = '润津实时出清结果'
    else:
        template = _get_template(int(mmdd[:2]), 'dayahead')
        out_dir = DA_OUT
        out_name = f'{mmdd}-日前出清收益测算.xlsx'
        sql = f"""SELECT date, time_point, power, price
FROM shandong_px_reliable_clearing_unit_data
WHERE date = '{date_iso}' AND member_id = 'b9e64e64a713458eba94c9af05c0a757'
  AND unit_name NOT LIKE '%发电%' AND unit_name NOT LIKE '%用电%'
ORDER BY time_point"""
        table_name = '日前出清收益测算'
        raw_table = '润津日前出清结果'

    out_path = os.path.join(out_dir, out_name)

    conn = get_conn()
    cur = conn.cursor()
    cur.execute(sql)
    rows = cur.fetchall()
    cur.close()
    conn.close()

    if not rows:
        print(f'  SKIP {mmdd} ({mode}): no data in 天机')
        return False

    # Extract 96-point data
    prices = [float(r[3]) if r[3] is not None else 0.0 for r in rows]
    powers = [float(r[2]) if r[2] is not None else 0.0 for r in rows]

    # 1. Copy template
    shutil.copy2(template, out_path)

    # 2. Write 报价及预中标
    wb = openpyxl.load_workbook(out_path)
    ws = wb['报价及预中标']

    for i in range(96):
        r = i + 2  # Row 2-97
        # H = timestamp
        h_cell = ws.cell(row=r, column=8)
        if not _is_merged(h_cell):
            h_cell.value = TIMES[i]
        # J = price
        j_cell = ws.cell(row=r, column=10)
        if not _is_merged(j_cell):
            j_cell.value = prices[i]
        # K = price (same as J for both modes)
        k_cell = ws.cell(row=r, column=11)
        if not _is_merged(k_cell):
            k_cell.value = prices[i]
        # N = power
        n_cell = ws.cell(row=r, column=14)
        if not _is_merged(n_cell):
            n_cell.value = powers[i]

    wb.save(out_path)
    wb.close()
    print(f'  [OK] {out_name}')

    # 3. COM refresh
    print(f'  COM refresh {mmdd} ({mode}) ...')
    com_refresh(out_path)

    # 4. Read Row 4 (17 cols: A-Q)
    wb2 = openpyxl.load_workbook(out_path, data_only=True)
    ws2 = wb2['充放测算']
    row4 = [ws2.cell(row=4, column=c).value for c in range(1, 18)]
    summary = [safe_float(v) for v in row4]
    wb2.close()

    # 5. Write to SQLite
    db = sqlite3.connect(DB_PATH)
    cur = db.cursor()

    cur.execute(f'PRAGMA table_info("{table_name}")')
    cols = [r[1] for r in cur.fetchall()]
    placeholders = ','.join(['?'] * len(cols))
    col_list = ','.join([f'"{c}"' for c in cols])

    cur.execute(f'DELETE FROM "{table_name}" WHERE date=?', (mmdd,))
    cur.execute(f'INSERT OR REPLACE INTO "{table_name}" ({col_list}) VALUES ({placeholders})',
                (mmdd, *summary))

    db.commit()
    db.close()
    print(f'  [DB] {table_name}: 1 row (17 cols)')

    return True


def main():
    # 天机有数据的实时出清日期
    realtime_dates = ['0708','0712','0713','0714','0715','0716','0717','0718',
                      '0719','0720','0721','0722','0723','0724','0725','0726',
                      '0727','0728','0729','0730','0731','0801','0802','0803','0804','0805']
    # 天机有数据的日前出清日期
    dayahead_dates = ['0711','0712','0713','0714','0715','0716','0717','0718',
                      '0719','0720','0721','0722','0723','0724','0725','0726',
                      '0727','0728','0729','0730','0731','0801','0802','0803','0804','0805','0806']

    db = sqlite3.connect(DB_PATH)
    cur = db.cursor()
    cur.execute("SELECT date FROM 实时出清收益测算")
    existing_rt = set(r[0] for r in cur.fetchall())
    cur.execute("SELECT date FROM 日前出清收益测算")
    existing_da = set(r[0] for r in cur.fetchall())
    db.close()

    todo_rt = [d for d in realtime_dates if d not in existing_rt]
    todo_da = [d for d in dayahead_dates if d not in existing_da]

    print(f'=== 实时出清收益测算 ===')
    print(f'天机有数据: {len(realtime_dates)} 天, 已入库: {len(existing_rt)} 天, 待生成: {len(todo_rt)}')
    print(f'=== 日前出清收益测算 ===')
    print(f'天机有数据: {len(dayahead_dates)} 天, 已入库: {len(existing_da)} 天, 待生成: {len(todo_da)}')

    for d in todo_rt:
        print(f'\n--- 实时 {d} ---')
        generate_one(d, mode='realtime')

    for d in todo_da:
        print(f'\n--- 日前 {d} ---')
        generate_one(d, mode='dayahead')

    print('\nDone.')


if __name__ == '__main__':
    main()